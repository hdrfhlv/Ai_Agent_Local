"""
scraper_tools.py - skill scraping + export Excel, independen dari provider LLM.
Install: pip install requests beautifulsoup4 pandas openpyxl
"""
import json
import math
import os
import re
import statistics
import urllib.robotparser
from urllib.parse import urljoin, urlparse

import pandas as pd
import requests
from bs4 import BeautifulSoup
from openpyxl.chart import BarChart, Reference
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

# ---------- Konfigurasi scraper dan header Excel ----------
HEADERS = {"User-Agent": "Mozilla/5.0 (compatible; MyScraperAgent/1.0)"}

EXCEL_HEADERS = {
    "nama": "Nama Produk",
    "harga": "Harga",
    "tautan": "Tautan Produk",
}


# ---------- Membuat tabel Excel yang sudah diformat ----------
def _save_excel_table(rows: list, filename: str, sheet_name: str = "Produk") -> None:
    df = pd.DataFrame(rows)
    df.rename(
        columns={
            column: EXCEL_HEADERS.get(str(column).strip().lower(), str(column).strip().title())
            for column in df.columns
        },
        inplace=True,
    )
    df.to_excel(filename, index=False, sheet_name=sheet_name)

    workbook = load_workbook(filename)
    worksheet = workbook[sheet_name]
    worksheet.freeze_panes = "A2"
    worksheet.row_dimensions[1].height = 24
    worksheet.sheet_view.showGridLines = False

    for column_cells in worksheet.columns:
        column_index = column_cells[0].column
        values = [str(cell.value or "") for cell in column_cells]
        width = min(max(max(map(len, values), default=0) + 2, 14), 55)
        worksheet.column_dimensions[get_column_letter(column_index)].width = width
        for cell in column_cells:
            cell.alignment = Alignment(vertical="center")
            if cell.row > 1 and isinstance(cell.value, str) and cell.value.startswith(("http://", "https://")):
                cell.hyperlink = cell.value
                cell.style = "Hyperlink"

    if worksheet.max_row >= 2 and worksheet.max_column >= 1:
        table = Table(
            displayName="ScrapedData",
            ref=f"A1:{get_column_letter(worksheet.max_column)}{worksheet.max_row}",
        )
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        worksheet.add_table(table)
    elif worksheet.max_column:
        worksheet.auto_filter.ref = worksheet.dimensions

    workbook.save(filename)


# ---------- Helper untuk mengambil dan mengekstrak halaman ----------
def _allowed(url: str) -> bool:
    p = urlparse(url)
    rp = urllib.robotparser.RobotFileParser()
    rp.set_url(f"{p.scheme}://{p.netloc}/robots.txt")
    try:
        rp.read()
        return rp.can_fetch(HEADERS["User-Agent"], url)
    except Exception:
        return True


def fetch_html(url: str, max_chars: int = 15000) -> str:
    if not _allowed(url):
        return "ERROR: dilarang oleh robots.txt"
    r = requests.get(url, headers=HEADERS, timeout=20)
    r.raise_for_status()
    soup = BeautifulSoup(r.text, "html.parser")
    for t in soup(["script", "style", "svg", "noscript"]):
        t.decompose()
    return soup.prettify()[:max_chars]


def extract_data(url: str, item_selector: str, fields: dict) -> list:
    if not _allowed(url):
        return [{"error": "dilarang oleh robots.txt"}]
    r = requests.get(url, headers=HEADERS, timeout=20)
    r.raise_for_status()
    soup = BeautifulSoup(r.text, "html.parser")
    rows = []
    for item in soup.select(item_selector):
        row = {}
        for name, spec in fields.items():
            sel, _, attr = spec.partition("@")
            el = item.select_one(sel) if sel else item
            row[name] = None if el is None else (el.get(attr) if attr else el.get_text(strip=True))
        rows.append(row)
    return rows


def export_to_excel(rows: list, filename: str = "hasil_scraping.xlsx") -> str:
    _save_excel_table(rows, filename)
    return f"Tersimpan: {filename} ({len(rows)} baris)"


# ---------- Membaca dan membersihkan nilai numerik ----------
def _parse_number(value) -> float | None:
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float)):
        number = float(value)
        return number if math.isfinite(number) else None

    match = re.search(r"-?\d[\d.,]*", str(value).replace("\u00a0", " "))
    if not match:
        return None
    number = match.group()
    if "," in number and "." in number:
        decimal_separator = "," if number.rfind(",") > number.rfind(".") else "."
        thousands_separator = "." if decimal_separator == "," else ","
        number = number.replace(thousands_separator, "").replace(decimal_separator, ".")
    elif "," in number or "." in number:
        separator = "," if "," in number else "."
        whole, _, fraction = number.rpartition(separator)
        if len(fraction) in (1, 2):
            number = f"{whole.replace(separator, '')}.{fraction}"
        else:
            number = number.replace(separator, "")
    try:
        parsed = float(number)
    except ValueError:
        return None
    return parsed if math.isfinite(parsed) else None


def _column_label(column) -> str:
    return str(column).strip().lower()


# ---------- Menganalisis data dan menambahkan diagram ke Excel ----------
def analyze_scraped_data(filename: str = "hasil_scraping.xlsx") -> dict:
    """Tambahkan ringkasan statistik dan diagram ke workbook hasil scraping."""
    if not os.path.isfile(filename):
        return {"ok": False, "error": f"File Excel tidak ditemukan: {filename}"}

    try:
        workbook = load_workbook(filename)
        source_sheet = "Produk" if "Produk" in workbook.sheetnames else workbook.sheetnames[0]
        data = pd.read_excel(filename, sheet_name=source_sheet)
    except (OSError, ValueError, KeyError) as exc:
        return {"ok": False, "error": f"Gagal membaca file Excel: {exc}"}

    data = data.dropna(how="all")
    if data.empty or len(data.columns) == 0:
        return {"ok": False, "error": "Lembar kerja tidak memiliki data untuk dianalisis."}

    numeric_columns = []
    parsed_columns = {}
    for column in data.columns:
        label = _column_label(column)
        if any(term in label for term in ("url", "tautan", "link", "kode", " id")):
            continue
        values = data[column].dropna().tolist()
        parsed = [_parse_number(value) for value in values]
        numbers = [number for number in parsed if number is not None]
        is_numeric_label = any(
            term in label
            for term in ("harga", "price", "jumlah", "quantity", "qty", "stok", "stock",
                         "rating", "nilai", "score", "total", "amount", "revenue")
        )
        if numbers and (is_numeric_label or len(numbers) / len(values) >= 0.8):
            numeric_columns.append(column)
            parsed_columns[column] = [
                _parse_number(value) for value in data[column].tolist()
            ]

    analysis_sheet_name = "Analisis"
    if analysis_sheet_name in workbook.sheetnames:
        del workbook[analysis_sheet_name]
    sheet = workbook.create_sheet(analysis_sheet_name)
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A7"
    sheet.merge_cells("A1:G1")
    sheet["A1"] = "Ringkasan Analisis Data"
    sheet["A1"].font = Font(bold=True, size=16, color="FFFFFF")
    sheet["A1"].fill = PatternFill("solid", fgColor="1F4E78")
    sheet["A1"].alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 30
    sheet["A2"] = "Sumber data"
    sheet["B2"] = source_sheet
    sheet["A3"] = "Jumlah baris data"
    sheet["B3"] = len(data)

    header_fill = PatternFill("solid", fgColor="4472C4")
    header_font = Font(bold=True, color="FFFFFF")
    summary = []
    if numeric_columns:
        sheet["A5"] = "Statistik numerik"
        summary_headers = ["Kolom", "Jumlah", "Rata-rata", "Median", "Minimum", "Maksimum"]
        for column_index, value in enumerate(summary_headers, start=1):
            cell = sheet.cell(row=6, column=column_index, value=value)
            cell.fill = header_fill
            cell.font = header_font

        for row_index, column in enumerate(numeric_columns, start=7):
            values = [value for value in parsed_columns[column] if value is not None]
            summary.append(
                {
                    "kolom": str(column),
                    "jumlah": len(values),
                    "rata_rata": statistics.mean(values),
                    "median": statistics.median(values),
                    "minimum": min(values),
                    "maksimum": max(values),
                }
            )
            sheet.cell(row=row_index, column=1, value=str(column))
            sheet.cell(row=row_index, column=2, value=len(values))
            sheet.cell(row=row_index, column=3, value=statistics.mean(values))
            sheet.cell(row=row_index, column=4, value=statistics.median(values))
            sheet.cell(row=row_index, column=5, value=min(values))
            sheet.cell(row=row_index, column=6, value=max(values))
            for column_index in range(3, 7):
                sheet.cell(row=row_index, column=column_index).number_format = "#,##0.00"

        chart_column = numeric_columns[0]
        chart_values = parsed_columns[chart_column]
        label_column = next(
            (
                column for column in data.columns
                if column != chart_column
                and any(term in _column_label(column) for term in ("nama", "name", "produk", "product"))
            ),
            next((column for column in data.columns if column != chart_column), None),
        )
        ranked = []
        for index, value in enumerate(chart_values):
            if value is None:
                continue
            label = data.iloc[index][label_column] if label_column is not None else f"Baris {index + 1}"
            ranked.append((str(label), value))
        ranked.sort(key=lambda item: item[1], reverse=True)
        chart_title = f"10 Nilai Tertinggi - {chart_column}"
        y_axis_title = str(chart_column)
        x_axis_title = "Kategori"
    else:
        label_column = next(
            (column for column in data.columns if data[column].notna().any()),
            None,
        )
        if label_column is None:
            return {"ok": False, "error": "Tidak ada kolom yang dapat dianalisis."}
        counts = data[label_column].dropna().astype(str).value_counts().head(10)
        if counts.empty:
            return {"ok": False, "error": "Tidak ada nilai untuk dibuat diagram."}
        sheet["A5"] = "Frekuensi kategori"
        for column_index, value in enumerate([str(label_column), "Jumlah"], start=1):
            cell = sheet.cell(row=6, column=column_index, value=value)
            cell.fill = header_fill
            cell.font = header_font
        ranked = [(str(label), int(count)) for label, count in counts.items()]
        summary.append(
            {
                "kolom": str(label_column),
                "kategori_unik": int(data[label_column].nunique()),
                "kategori_teratas": [
                    {"kategori": label, "jumlah": count}
                    for label, count in ranked[:5]
                ],
            }
        )
        chart_title = f"10 Kategori Teratas - {label_column}"
        y_axis_title = "Jumlah"
        x_axis_title = str(label_column)

    sheet["D5"] = "Data untuk diagram"
    for column_index, value in enumerate([x_axis_title, y_axis_title], start=4):
        cell = sheet.cell(row=6, column=column_index, value=value)
        cell.fill = header_fill
        cell.font = header_font
    for row_index, (label, value) in enumerate(ranked[:10], start=7):
        sheet.cell(row=row_index, column=4, value=label)
        value_cell = sheet.cell(row=row_index, column=5, value=value)
        if numeric_columns:
            value_cell.number_format = "#,##0.00"

    chart = BarChart()
    chart.type = "bar"
    chart.style = 10
    chart.title = chart_title
    chart.y_axis.title = y_axis_title
    chart.x_axis.title = x_axis_title
    chart.height = 9
    chart.width = 18
    chart.add_data(
        Reference(sheet, min_col=5, min_row=6, max_row=6 + min(len(ranked), 10)),
        titles_from_data=True,
    )
    chart.set_categories(
        Reference(sheet, min_col=4, min_row=7, max_row=6 + min(len(ranked), 10))
    )
    sheet.add_chart(chart, "G5")

    sheet.column_dimensions["A"].width = 24
    sheet.column_dimensions["B"].width = 18
    sheet.column_dimensions["C"].width = 16
    sheet.column_dimensions["D"].width = 36
    sheet.column_dimensions["E"].width = 18
    sheet.column_dimensions["F"].width = 18
    try:
        workbook.save(filename)
    except (OSError, ValueError) as exc:
        return {"ok": False, "error": f"Gagal menyimpan analisis ke Excel: {exc}"}

    return {
        "ok": True,
        "file": filename,
        "lembar_analisis": analysis_sheet_name,
        "jumlah_baris": len(data),
        "kolom_numerik": [str(column) for column in numeric_columns],
        "diagram": chart_title,
        "ringkasan": summary,
        "data_diagram": [
            {"kategori": label, "nilai": value}
            for label, value in ranked[:5]
        ],
    }


# ---------- Spesifikasi tool dalam format JSON Schema ----------
TOOL_SPECS = [
    {
        "name": "scrape_products",
        "description": "Ambil nama, harga jika tersedia, dan tautan produk dari katalog publik. Ikuti halaman berikutnya otomatis sampai max_items tercapai atau halaman habis, lalu simpan data dan analisis ke Excel. filename opsional untuk menentukan nama file .xlsx.",
        "parameters": {
            "type": "object",
            "properties": {
                "url": {"type": "string", "description": "URL lengkap katalog atau produk."},
                "max_items": {"type": "integer", "description": "Jumlah maksimum produk yang diambil lintas halaman (default 50)."},
                "filename": {"type": "string", "description": "Nama file Excel keluaran, opsional; default hasil_scraping.xlsx. Hanya nama file, tanpa path."},
            },
            "required": ["url"],
        },
    },
    {
        "name": "analyze_scraped_data",
        "description": "Analisis file Excel hasil scraping. Tambahkan lembar Analisis dengan statistik dan diagram batang, tanpa mengubah data sumber. Jika nama file tidak disebut, gunakan hasil_scraping.xlsx.",
        "parameters": {
            "type": "object",
            "properties": {
                "filename": {"type": "string", "description": "Lokasi file .xlsx yang akan dianalisis."},
            },
        },
    },
    {
        "name": "fetch_html",
        "description": "Ambil HTML halaman untuk menganalisis struktur dan menentukan CSS selector.",
        "parameters": {
            "type": "object",
            "properties": {"url": {"type": "string"}},
            "required": ["url"],
        },
    },
    {
        "name": "extract_data",
        "description": "Ekstrak data berulang dari halaman memakai CSS selector. Gunakan 'selector@atribut' untuk atribut seperti href.",
        "parameters": {
            "type": "object",
            "properties": {
                "url": {"type": "string"},
                "item_selector": {"type": "string"},
                "fields": {"type": "object", "additionalProperties": {"type": "string"}},
            },
            "required": ["url", "item_selector", "fields"],
        },
    },
    {
        "name": "export_to_excel",
        "description": "Simpan list of dict ke file .xlsx.",
        "parameters": {
            "type": "object",
            "properties": {
                "rows": {"type": "array", "items": {"type": "object"}},
                "filename": {"type": "string"},
            },
            "required": ["rows"],
        },
    },
]


# ---------- Mengambil data produk dari JSON-LD ----------
def _product_from_jsonld(data, page_url: str) -> list:
    products = []

    def visit(value):
        if isinstance(value, list):
            for item in value:
                visit(item)
            return
        if not isinstance(value, dict):
            return

        product_type = value.get("@type", [])
        if isinstance(product_type, str):
            product_type = [product_type]
        if any("product" in str(item).lower() for item in product_type):
            offers = value.get("offers") or {}
            if isinstance(offers, list):
                offers = offers[0] if offers else {}
            if not isinstance(offers, dict):
                offers = {}
            price = offers.get("price") or offers.get("lowPrice")
            name = value.get("name")
            if name:
                name = str(name).strip()
                products.append(
                    {
                        "nama": name,
                        "harga": price,
                        "tautan": urljoin(page_url, value.get("url") or page_url),
                    }
                )
        for nested in value.values():
            if isinstance(nested, (dict, list)):
                visit(nested)

    visit(data)
    return products


# ---------- Mengambil data produk dari elemen HTML ----------
def _products_from_html(soup, page_url: str) -> list:
    products = []
    cards = soup.select(
        '[itemscope][itemtype*="Product"], .product, '
        '[class*="product-card"], [class*="product-item"], [data-product], '
        'article.product_pod'
    )
    for card in cards:
        name_el = card.select_one(
            '[itemprop="name"], .product-title, .product-name'
        )
        if name_el is None:
            name_el = card.select_one("h2 a[title], h3 a[title]")
        if name_el is None:
            name_el = card.select_one("h2, h3")
        price_el = card.select_one('[itemprop="price"], .price_color, .price')
        if price_el is None:
            price_el = card.select_one('[class*="price"]')
        link_el = card.select_one('a[itemprop="url"], a[href]')
        name = (
            name_el.get("content")
            or name_el.get("title")
            or name_el.get_text(" ", strip=True)
            if name_el
            else None
        )
        price = (
            price_el.get("content") or price_el.get_text(" ", strip=True)
            if price_el
            else None
        )
        link = urljoin(page_url, link_el.get("href")) if link_el else page_url
        if name:
            products.append({"nama": name, "harga": price, "tautan": link})
    return products


# ---------- Menggabungkan data produk dari satu halaman ----------
def _products_from_page(soup, page_url: str) -> list:
    products = []
    for script in soup.select('script[type="application/ld+json"]'):
        try:
            data = json.loads(script.string or script.get_text())
        except (json.JSONDecodeError, TypeError):
            continue
        products.extend(_product_from_jsonld(data, page_url))
    products.extend(_products_from_html(soup, page_url))
    return products


# ---------- Menemukan tautan ke halaman katalog berikutnya ----------
def _next_page_url(soup, page_url: str) -> str | None:
    next_link = soup.select_one(
        'a[rel~="next"], a[aria-label*="next" i], '
        ".next a[href], a.next[href], li.next a[href]"
    )
    if next_link is None:
        for link in soup.select("a[href]"):
            label = link.get_text(" ", strip=True).lower()
            if label == "next" or label.startswith("next "):
                next_link = link
                break
    if next_link is None:
        return None
    href = next_link.get("href")
    return urljoin(page_url, href) if href else None


# ---------- Scrape katalog lintas halaman dan ekspor hasil ----------
def scrape_products(
    url: str,
    max_items: int = 50,
    filename: str = "hasil_scraping.xlsx",
) -> dict:
    """Ambil produk lintas halaman katalog sampai batas item lalu ekspor ke Excel."""
    if not isinstance(url, str) or not url.startswith(("http://", "https://")):
        return {"ok": False, "error": "URL harus diawali http:// atau https://"}
    if isinstance(max_items, bool) or not isinstance(max_items, int) or max_items < 1:
        return {"ok": False, "error": "max_items harus lebih besar dari 0"}
    if not isinstance(filename, str) or not filename.strip():
        return {"ok": False, "error": "Nama file Excel tidak boleh kosong"}
    filename = filename.strip()
    if any(separator in filename for separator in ("/", "\\", ":")) or filename in {".", ".."}:
        return {"ok": False, "error": "Masukkan nama file saja, tanpa path folder"}
    extension = os.path.splitext(filename)[1]
    if extension and extension.lower() != ".xlsx":
        return {"ok": False, "error": "File keluaran harus menggunakan format .xlsx"}
    if not extension:
        filename = f"{filename}.xlsx"

    current_url = url
    visited_pages = set()
    unique_products = []
    seen = set()
    page_error = None
    while current_url and current_url not in visited_pages and len(unique_products) < max_items:
        if not _allowed(current_url):
            if not visited_pages:
                return {"ok": False, "error": "Scraping dilarang oleh robots.txt"}
            page_error = f"Halaman berikutnya diblokir robots.txt: {current_url}"
            break
        visited_pages.add(current_url)

        try:
            response = requests.get(current_url, headers=HEADERS, timeout=20)
            response.raise_for_status()
        except requests.RequestException as exc:
            if not unique_products:
                return {"ok": False, "error": str(exc)}
            page_error = f"Gagal mengambil halaman berikutnya {current_url}: {exc}"
            break

        page_url = response.url
        soup = BeautifulSoup(response.text, "html.parser")
        for product in _products_from_page(soup, page_url):
            key = (str(product["nama"]).strip().lower(), product["tautan"])
            if key not in seen:
                seen.add(key)
                unique_products.append(product)
            if len(unique_products) >= max_items:
                break
        current_url = _next_page_url(soup, page_url)

    if not unique_products:
        return {"ok": False, "error": "Tidak menemukan data produk pada halaman tersebut."}

    detail_errors = []
    for product in unique_products:
        detail_url = product["tautan"]
        if product["harga"] is not None and str(product["harga"]).strip():
            continue
        if detail_url == url:
            continue
        if not _allowed(detail_url):
            detail_errors.append(f"Detail produk diblokir robots.txt: {detail_url}")
            continue
        try:
            detail_response = requests.get(detail_url, headers=HEADERS, timeout=20)
            detail_response.raise_for_status()
        except requests.RequestException as exc:
            detail_errors.append(f"Gagal membaca detail {detail_url}: {exc}")
            continue

        detail_soup = BeautifulSoup(detail_response.text, "html.parser")
        detail_products = _products_from_page(detail_soup, detail_url)
        if detail_products:
            detail_product = next(
                (
                    item for item in detail_products
                    if item["nama"]
                    and str(item["nama"]).strip().lower()
                    == str(product["nama"]).strip().lower()
                ),
                detail_products[0],
            )
            product["harga"] = detail_product["harga"]

    try:
        _save_excel_table(unique_products, filename)
    except (OSError, ValueError, ImportError) as exc:
        return {"ok": False, "error": f"Gagal menyimpan file Excel: {exc}"}

    analysis = analyze_scraped_data(filename)
    result = {
        "ok": True,
        "jumlah": len(unique_products),
        "halaman": len(visited_pages),
        "file": filename,
    }
    if analysis.get("ok"):
        result["lembar_analisis"] = analysis["lembar_analisis"]
        result["diagram"] = analysis["diagram"]
        result["ringkasan"] = analysis["ringkasan"]
        result["data_diagram"] = analysis["data_diagram"]
    else:
        result["catatan"] = (
            f"Data scrape tersimpan, tetapi analisis otomatis gagal: "
            f"{analysis.get('error', 'alasan tidak diketahui')}"
        )
    if page_error:
        result["catatan"] = (
            f"{result['catatan']}; {page_error}"
            if result.get("catatan")
            else page_error
        )
    if detail_errors:
        detail_note = "; ".join(detail_errors)
        result["catatan"] = f"{result['catatan']}; {detail_note}" if result.get("catatan") else detail_note
    return result


# ---------- Mendaftarkan fungsi dan skema tool ----------
_FUNCS = {
    f.__name__: f
    for f in (scrape_products, analyze_scraped_data, fetch_html, extract_data, export_to_excel)
}

_PRODUCT_TOOL_SPEC = TOOL_SPECS[0]
TOOL_SCHEMA = {
    "name": _PRODUCT_TOOL_SPEC["name"],
    "description": _PRODUCT_TOOL_SPEC["description"],
    "input_schema": _PRODUCT_TOOL_SPEC["parameters"],
}
ACTIVE_TOOL_SCHEMAS = [TOOL_SCHEMA, {
    "name": TOOL_SPECS[1]["name"],
    "description": TOOL_SPECS[1]["description"],
    "input_schema": TOOL_SPECS[1]["parameters"],
}]


# ---------- Menjalankan tool berdasarkan nama ----------
def run_tool(name: str, args: dict) -> str:
    """Jalankan tool dan selalu kembalikan string (aman dikirim balik ke LLM)."""
    try:
        out = _FUNCS[name](**args)
    except Exception as e:
        out = f"ERROR: {e}"
    return out if isinstance(out, str) else json.dumps(out, ensure_ascii=False)[:20000]


# ---------- Konverter format per provider ----------
def as_anthropic_tools():
    return [{"name": t["name"], "description": t["description"], "input_schema": t["parameters"]}
            for t in TOOL_SPECS]


def as_openai_tools():  # juga dipakai provider yang OpenAI-compatible (Groq, DeepSeek, Ollama, dll)
    return [{"type": "function", "function": t} for t in TOOL_SPECS]